import sys
import os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../../')))
import os
import pandas as pd
import docx
from rag import ingest_documents, query_rag
import json

# 1. Create test files
os.makedirs("test_eval_docs", exist_ok=True)

# DOCX
doc = docx.Document()
doc.add_paragraph("This is a standard paragraph in the document.")
table = doc.add_table(rows=2, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Server Name'
hdr_cells[1].text = 'IP Address'
hdr_cells[2].text = 'Status'
row_cells = table.rows[1].cells
row_cells[0].text = 'Alpha-DB-1'
row_cells[1].text = '192.168.1.100'
row_cells[2].text = 'Critical Failure'
doc.save("test_eval_docs/test_table.docx")

# XLSX
df = pd.DataFrame({
    'Requirement ID': ['REQ-101', 'REQ-102'],
    'Description': ['System must encrypt data at rest', 'System must log all access attempts'],
    'Status': ['Implemented', 'Pending']
})
df.to_excel("test_eval_docs/test_mapping.xlsx", index=False)

# 2. Ingest
ingest_documents("test_eval_docs")

# 3. Query
queries = {
    "docx_table": "What is the status of the Alpha-DB-1 server?",
    "xlsx_exact": "What is the description for REQ-102?",
    "xlsx_para": "Which requirement mandates keeping a record of login attempts?"
}

print("=== EVALUATION RESULTS ===")
for key, q in queries.items():
    res = query_rag(q, top_k=1)
    if res["context"]:
        print(f"{key}_score:", res["context"][0]["score"])
        print(f"{key}_text:", repr(res["context"][0]["text"][:100]))
    else:
        print(f"{key}_score: 0.0")

