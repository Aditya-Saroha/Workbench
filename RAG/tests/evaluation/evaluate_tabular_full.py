import sys
import os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../../')))
import os
import pandas as pd
import docx
from PIL import Image
from rag import ingest_documents, query_rag
import json

os.makedirs("test_eval_docs", exist_ok=True)

# 1. DOCX (Paragraph + Table)
doc = docx.Document()
doc.add_paragraph("This is a standard paragraph in the document containing instructions for the Delta protocol.")
table = doc.add_table(rows=2, cols=3)
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = 'Server Name', 'IP Address', 'Status'
row = table.rows[1].cells
row[0].text, row[1].text, row[2].text = 'Alpha-DB-1', '192.168.1.100', 'Critical Failure'
doc.save("test_eval_docs/test_doc.docx")

# 2. XLSX (Multi-sheet)
df1 = pd.DataFrame({'Requirement ID': ['REQ-101', 'REQ-102'], 'Description': ['System must encrypt data at rest', 'System must log all access attempts'], 'Status': ['Implemented', 'Pending']})
df2 = pd.DataFrame({'Asset': ['Firewall', 'Router'], 'Owner': ['Alice', 'Bob']})
with pd.ExcelWriter('test_eval_docs/test_xl.xlsx') as writer:
    df1.to_excel(writer, sheet_name='Requirements', index=False)
    df2.to_excel(writer, sheet_name='Assets', index=False)

# 3. Scanned PDF
img = Image.new("RGB", (800, 200), color="white")
from PIL import ImageDraw
d = ImageDraw.Draw(img)
d.text((10, 10), "CONFIDENTIAL SCANNED PASSCODE: TANGO-11", fill="black")
img.save("test_eval_docs/test_scanned.pdf", "PDF", resolution=100.0)

# 4. Normal PDF (We'll use a tiny synthetic one for speed, or copy one)
import pymupdf
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((50, 50), "NIST SP 800-172 Supplement: Enhanced Security Requirements.")
pdf.save("test_eval_docs/test_normal.pdf")

ingest_documents("test_eval_docs")

queries = {
    "A. DOCX Paragraph": "What protocol instructions are in the document?",
    "B. DOCX Table": "What is the status of the Alpha-DB-1 server?",
    "C. XLSX Exact": "What is the description for REQ-102?",
    "D. XLSX Paraphrased": "Which requirement mandates keeping a record of login attempts?",
    "E. XLSX Multi-sheet": "Who is the owner of the Firewall asset?",
    "F. Normal PDF": "What does the supplement cover?",
    "G. Scanned OCR": "What is the confidential scanned passcode?",
    "H. Out-of-Domain": "What is the recipe for chocolate cake?"
}

print("=== EVALUATION RESULTS ===")
for key, q in queries.items():
    res = query_rag(q, top_k=1)
    if res["context"]:
        print(f"{key} | Score: {res['context'][0]['score']:.4f} | Text: {repr(res['context'][0]['text'][:80])}")
    else:
        print(f"{key} | No Result")
