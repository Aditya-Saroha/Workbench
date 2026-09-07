import json
import traceback
import sys
import os
import urllib.request
import urllib.error
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt

# Add RAG folder to path
root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
rag_dir = os.path.join(root_dir, "RAG")
sys.path.append(rag_dir)

try:
    from rag import query_rag
except ImportError as e:
    print(f"Warning: Could not import query_rag. Error: {e}")
    query_rag = None

# agent.py and main.py both import this name — it went missing when
# direct_chat's model got hardcoded inline below, which is what sent an
# IDE auto-import hunting through the rest of the workspace and landing on
# a stale copy in save_from_git/.
DIRECT_CHAT_MODEL = "qwen3:4b"
DIRECT_CHAT_ROUTE = "simple_chat"

PROJECT_ROOT = Path(root_dir)
SESSION_OUTPUTS = PROJECT_ROOT / "RAG" / "sessions"


def _safe_read_path(filepath: str, session_id: str) -> Path:
    requested = Path(filepath)
    candidates = []
    if not requested.is_absolute():
        candidates.append(PROJECT_ROOT / requested)
        candidates.append(SESSION_OUTPUTS / _safe_session_id(session_id) / "documents" / requested.name)
    else:
        candidates.append(requested)

    blocked_parts = {".git", ".venv", "node_modules", "__pycache__"}
    for candidate in candidates:
        resolved = candidate.resolve()
        if any(part in blocked_parts for part in resolved.parts):
            continue
        try:
            resolved.relative_to(PROJECT_ROOT.resolve())
        except ValueError:
            continue
        if resolved.is_file():
            return resolved

    raise FileNotFoundError(f"File not found in the allowed workspace: {filepath}")


def _safe_session_id(session_id: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]", "-", session_id or "default")[:80] or "default"

# Add project root for sandbox package (root_dir already computed above)
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)

try:
    from sandbox.executor import run_sandboxed_session, close_sandbox_session, reap_idle_sandbox_sessions, TIMEOUT_SECONDS
except ImportError as e:
    print(f"Warning: Could not import sandbox executor. Error: {e}")
    run_sandboxed_session = None
    close_sandbox_session = None
    reap_idle_sandbox_sessions = None
    TIMEOUT_SECONDS = 10


def execute_tool(tool_name: str, tool_args: dict, session_id: str = "default") -> str:
    """Dispatcher for all agent tools. Handles validation, size limits, and robust errors."""
    
    # 1. Structural Validation
    if not isinstance(tool_args, dict):
        return f"Error: tool_args must be a dictionary, got {type(tool_args).__name__}"

    # 2. Strict Tool Routing & Argument Validation
    if tool_name == "direct_chat":
        prompt = tool_args.get("prompt")
        if not isinstance(prompt, str) or not prompt.strip():
            return "Error: No valid 'prompt' provided to direct_chat tool."
            
        # ollama-agent-router routes by "router": {"taskType": ...}, not by
        # putting the task type in "model" (that field wants "auto" or a
        # real Ollama model tag). "mode": "sync" + "allowAsync": False keeps
        # this from ever getting back an async job envelope instead of a
        # normal choices response.
        payload = json.dumps({
            "model": "auto",
            "messages": [{"role": "user", "content": prompt}],
            "router": {"taskType": DIRECT_CHAT_ROUTE, "mode": "sync", "allowAsync": False},
        }).encode("utf-8")
        req = urllib.request.Request("http://127.0.0.1:11435/v1/chat/completions", data=payload, headers={"Content-Type": "application/json"})
        
        try:
            with urllib.request.urlopen(req, timeout=300) as response:
                result = json.loads(response.read().decode("utf-8"))
                return result.get("choices", [{}])[0].get("message", {}).get("content", "No response generated.")
        except Exception as e:
            return f"Error: Failed to reach Ollama via direct_chat: {str(e)}"
            
    elif tool_name == "search_rag":
        query = tool_args.get("query")
        if not isinstance(query, str) or not query.strip():
            return "Error: No valid 'query' string provided to search_rag tool."
            
        if not query_rag: 
            return "Error: RAG search is unavailable. Start the RAG service or check its import dependencies."
            
        try:
            result = query_rag(query, top_k=5, session_id=session_id)
            contexts = result.get("context", [])
            if not contexts: 
                return f"No relevant documents found for query: '{query}'"
                
            formatted_context = [f"Found {len(contexts)} relevant document excerpts:\n"]
            for i, chunk in enumerate(contexts):
                text = chunk.get("text", "")
                source = chunk.get("source", "Unknown Document")
                page = chunk.get("page", "Unknown Page")
                formatted_context.append(f"--- Excerpt {i+1} ---\nSource: {source} (Page: {page})\nContent: {text.strip()}\n")
                
            out = "\n".join(formatted_context)
            # Size limit protection
            if len(out) > 500_000:
                out = out[:500_000] + "\n... [TRUNCATED - RAG output exceeded 500KB]"
            return out
        except Exception as e:
            return f"Error: executing RAG search failed: {str(e)}"
            
    elif tool_name == "read_file":
        filepath = tool_args.get("filepath")
        if not isinstance(filepath, str) or not filepath.strip():
            return "Error: No valid 'filepath' string provided to read_file tool."
        try:
            path = _safe_read_path(filepath, session_id)
            content = path.read_text(encoding="utf-8", errors="replace")
            if len(content) > 500_000:
                content = content[:500_000] + "\n... [TRUNCATED at 500KB]"
            return f"File: {path.relative_to(PROJECT_ROOT)}\n\n{content}"
        except Exception as e:
            return f"Error: could not read file: {e}"
        
    elif tool_name == "python_sandbox":
        # Accept either a single `code` string (backward compatible with
        # existing plans) or a `commands` list, so the planner can run
        # several statements one after another in the SAME sandbox
        # container for this chat — variables and imports persist across
        # commands, and across separate python_sandbox steps/calls within
        # the same session_id, instead of each call getting a throwaway
        # interpreter.
        commands = tool_args.get("commands")
        if commands is None:
            code = tool_args.get("code")
            if code is None:
                return "Error: Missing 'code' or 'commands' argument for python_sandbox."
            if not isinstance(code, str):
                return "Error: 'code' argument must be a string."
            commands = [code]
        elif not isinstance(commands, list) or not all(isinstance(c, str) for c in commands):
            return "Error: 'commands' must be a list of code strings."

        commands = [c for c in commands if c.strip()]
        if not commands:
            return "Error: 'code'/'commands' argument cannot be empty."

        if not run_sandboxed_session:
            return "Error: python_sandbox is unavailable — sandbox package could not be imported."

        try:
            entries = run_sandboxed_session(session_id, commands)
        except Exception as e:
            return f"Error: Sandbox execution failed unexpectedly: {str(e)}"

        # Two representations of the same run in one string:
        #   1. A plain "$ command" / output transcript, so the critic
        #      model (evaluate_step) and any plain-text consumer can read
        #      it directly, same as before.
        #   2. A machine-readable JSON block the chat UI parses to render
        #      a real scrollback terminal (each command as its own entry)
        #      instead of one flat blob.
        transcript_lines = []
        for entry in entries:
            if entry.get("command") is not None:
                transcript_lines.append(f"$ {entry['command'].strip()}")
            output_text = (entry.get("output") or "").rstrip()
            transcript_lines.append(output_text if output_text else "(no output)")
            if entry.get("timed_out"):
                transcript_lines.append(f"[timed out after {TIMEOUT_SECONDS}s — sandbox session was reset]")
            elif entry.get("exit_code") is not None:
                transcript_lines.append(f"[exit code: {entry['exit_code']}]")
            transcript_lines.append("")
        transcript = "\n".join(transcript_lines).rstrip()

        payload = json.dumps({"kind": "sandbox_transcript", "commands": entries})
        return f"{transcript}\n\n<!--SANDBOX_JSON-->{payload}<!--/SANDBOX_JSON-->"
        
    elif tool_name == "write_docx":
        filename = tool_args.get("filename")
        if not isinstance(filename, str) or not filename.strip():
            return "Error: No valid 'filename' string provided to write_docx tool."
        safe_name = Path(filename).name
        if not safe_name.lower().endswith(".docx"):
            safe_name += ".docx"
        output_dir = SESSION_OUTPUTS / _safe_session_id(session_id) / "outputs"
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / safe_name

        document = Document()
        title = tool_args.get("title")
        content = tool_args.get("content", tool_args.get("body", ""))
        if title:
            document.add_heading(str(title), level=1)
        for paragraph in str(content).split("\n\n"):
            document.add_paragraph(paragraph.strip())
        document.save(output_path)
        return f"DOCX created: {output_path.relative_to(PROJECT_ROOT)}"
        
    else:
        # Unknown tool (malformed LLM output or hallucination)
        return f"Error: Unknown tool: {tool_name}"